# -*- coding: utf-8 -*-
"""Compile handbook/course/module_*.md (teaching-course markdown) into one
styled DOCX: VoxShield_Complete_Course.docx. Parses the STYLE.md contract:
# module, ## chapter, ### sub, > callout, LEARN: box, - / 1. lists,
**bold**, and | pipe | tables."""
import os, re, glob
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE=os.path.dirname(os.path.abspath(__file__))
COURSE=os.path.join(HERE,"course")
BLUE=RGBColor(0x1F,0x4E,0x8C); DARK=RGBColor(0x1A,0x1A,0x1A); GREY=RGBColor(0x55,0x55,0x55)
GREEN=RGBColor(0x1F,0x7A,0x33); ACCENT=RGBColor(0x2F,0x6F,0xE0); AMBER=RGBColor(0x8A,0x54,0x00)
CALLBG="EAF3EA"; LEARNBG="EAF1FB"

doc=Document()
st=doc.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(11); st.font.color.rgb=DARK
st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.12
for s in doc.sections:
    s.left_margin=s.right_margin=Inches(0.9); s.top_margin=s.bottom_margin=Inches(0.8)

def shade(p,hexcol):
    pr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),hexcol); pr.append(sh)
def border_left(p,hexcol):
    pr=p._p.get_or_add_pPr(); pb=OxmlElement('w:pBdr'); l=OxmlElement('w:left')
    l.set(qn('w:val'),'single'); l.set(qn('w:sz'),'18'); l.set(qn('w:space'),'8'); l.set(qn('w:color'),hexcol)
    pb.append(l); pr.append(pb)

def add_runs(p,text,base_color=DARK,size=11,italic=False):
    """render inline **bold** within text."""
    for i,seg in enumerate(re.split(r'(\*\*.+?\*\*)',text)):
        if not seg: continue
        b=seg.startswith("**") and seg.endswith("**")
        r=p.add_run(seg[2:-2] if b else seg); r.bold=b; r.italic=italic
        r.font.size=Pt(size); r.font.color.rgb=base_color

def H1(t):
    doc.add_page_break()
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(10)
    r=p.add_run(t); r.bold=True; r.font.size=Pt(22); r.font.color.rgb=BLUE
    bar=doc.add_paragraph(); bar.paragraph_format.space_after=Pt(10)
    pr=bar._p.get_or_add_pPr(); pb=OxmlElement('w:pBdr'); bot=OxmlElement('w:bottom')
    bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),'12'); bot.set(qn('w:space'),'1'); bot.set(qn('w:color'),'1F4E8C')
    pb.append(bot); pr.append(pb)
def H2(t):
    p=doc.add_heading(level=2); p.paragraph_format.space_before=Pt(14); p.paragraph_format.space_after=Pt(4)
    r=p.add_run(t); r.bold=True; r.font.size=Pt(14.5); r.font.color.rgb=ACCENT
def H3(t):
    p=doc.add_heading(level=3); p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(t); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=DARK

def callout(lines):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.1); p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(8)
    p.paragraph_format.space_after=Pt(8); shade(p,CALLBG); border_left(p,"1F7A33")
    r=p.add_run("In simple words:  "); r.bold=True; r.font.color.rgb=GREEN; r.font.size=Pt(11)
    add_runs(p," ".join(lines),base_color=RGBColor(0x24,0x4A,0x24),size=11,italic=True)
def learn_box(items):
    p=doc.add_paragraph(); shade(p,LEARNBG); border_left(p,"2F6FE0")
    p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(2)
    r=p.add_run("What you'll learn"); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=BLUE
    for it in items:
        b=doc.add_paragraph(style="List Bullet"); b.paragraph_format.space_after=Pt(1)
        add_runs(b,it,size=10.5,base_color=RGBColor(0x22,0x3A,0x5E))

def table(rows):
    hdr=rows[0]; t=doc.add_table(rows=1,cols=len(hdr)); t.style="Light Grid Accent 1"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(hdr):
        c=t.rows[0].cells[i].paragraphs[0]; add_runs(c,h,base_color=BLUE,size=9.5);
        for r in c.runs: r.bold=True
    for row in rows[1:]:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            if i<len(cells): add_runs(cells[i].paragraphs[0],v,size=9.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def para(text):
    p=doc.add_paragraph(); add_runs(p,text); return p
def bullet(text):
    p=doc.add_paragraph(style="List Bullet"); add_runs(p,text)
def numbered(text):
    p=doc.add_paragraph(style="List Number"); add_runs(p,text)

def is_table_row(l): return l.strip().startswith("|") and l.strip().endswith("|") and l.count("|")>=2
def parse_row(l): return [c.strip() for c in l.strip().strip("|").split("|")]
def is_sep(cells): return all(re.fullmatch(r':?-{2,}:?',c.strip()) for c in cells if c.strip())

# ---------- COVER ----------
def cover():
    for _ in range(2): doc.add_paragraph()
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("VoxShield"); r.bold=True; r.font.size=Pt(40); r.font.color.rgb=BLUE
    for txt,sz,col,it in [("The Complete Course",22,DARK,False),
      ("Understanding AI Voice-Clone Detection for Indian Banking — from the very beginning.",12,GREY,True),
      ("",6,DARK,False),
      ("A plain-English, step-by-step explanation of every part of the system —",11,DARK,False),
      ("the threat, the science, the models, and the live product.",11,DARK,False),
      ("",8,DARK,False),
      ("Team DigiSeva · PS2 · Live at 64.177.121.208.sslip.io",11,GREEN,False)]:
        q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if txt: rr=q.add_run(txt); rr.font.size=Pt(sz); rr.font.color.rgb=col; rr.italic=it
    for _ in range(2): doc.add_paragraph()
    box=doc.add_paragraph(); box.alignment=WD_ALIGN_PARAGRAPH.CENTER; shade(box,LEARNBG)
    r=box.add_run("How to read this book"); r.bold=True; r.font.color.rgb=BLUE; r.font.size=Pt(11)
    g=doc.add_paragraph(); g.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_runs(g,"Twelve short modules, in order. Each starts by explaining an idea in the simplest possible way, then builds up the detail. Green boxes give the one-line \"in simple words\" version. Every number is from the real, live system — nothing is exaggerated.",size=10.5,italic=True,base_color=GREY)

# ---------- PARSE ONE FILE ----------
def render_file(path):
    lines=open(path,encoding="utf-8").read().split("\n")
    i=0; para_buf=[]
    def flush():
        nonlocal para_buf
        if para_buf: para(" ".join(para_buf).strip()); para_buf=[]
    while i<len(lines):
        raw=lines[i]; l=raw.rstrip()
        s=l.strip()
        # table block
        if is_table_row(l):
            flush(); block=[]
            while i<len(lines) and is_table_row(lines[i]):
                cells=parse_row(lines[i])
                if not is_sep(cells): block.append(cells)
                i+=1
            if block: table(block)
            continue
        if s.startswith("# "): flush(); H1(s[2:].strip()); i+=1; continue
        if s.startswith("## "): flush(); H2(s[3:].strip()); i+=1; continue
        if s.startswith("### "): flush(); H3(s[4:].strip()); i+=1; continue
        if s.startswith(">"):
            flush(); block=[]
            while i<len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip(">").strip()); i+=1
            callout(block); continue
        if s.startswith("LEARN:"):
            flush(); i+=1; items=[]
            while i<len(lines) and lines[i].strip().startswith("- "):
                items.append(lines[i].strip()[2:].strip()); i+=1
            learn_box(items); continue
        if s.startswith("- "): flush(); bullet(s[2:].strip()); i+=1; continue
        if re.match(r'^\d+\.\s',s): flush(); numbered(re.sub(r'^\d+\.\s','',s)); i+=1; continue
        if s=="" : flush(); i+=1; continue
        if set(s)<=set("-—–") and len(s)>=3: i+=1; continue  # skip divider
        para_buf.append(s); i+=1
    flush()

# ---------- BUILD ----------
cover()
files=sorted(glob.glob(os.path.join(COURSE,"module_*.md")))
assert files, "no module_*.md found"
for f in files:
    render_file(f)
    print("rendered",os.path.basename(f))
out=os.path.join(HERE,"..","VoxShield_Complete_Course.docx")
doc.save(out); print("SAVED",os.path.abspath(out),"| modules:",len(files))
