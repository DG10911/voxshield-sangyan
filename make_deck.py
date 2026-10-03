"""Generate SANGYAN_VoxShield.pptx — 10-slide investor-resilience pitch deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BG = RGBColor(0x0B, 0x12, 0x20); ACC = RGBColor(0x22, 0xD3, 0xEE)
FG = RGBColor(0xE6, 0xED, 0xF3); SUB = RGBColor(0x94, 0xA3, 0xB8); WARN = RGBColor(0xF5, 0x9E, 0x0B)

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _bg(s):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG


def add(slide, title, bullets=None, kicker=None, accent=ACC):
    _bg(slide)
    if kicker:
        tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.5))
        p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = kicker.upper()
        r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = accent
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.9), Inches(12), Inches(1.0))
    p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = FG
    if bullets:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.7), Inches(4.9))
        tf = box.text_frame; tf.word_wrap = True
        for i, b in enumerate(bullets):
            lvl = 0; txt = b
            if isinstance(b, tuple): txt, lvl = b
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.level = lvl
            para.space_after = Pt(10)
            run = para.add_run(); run.text = ("• " if lvl == 0 else "– ") + txt
            run.font.size = Pt(19 if lvl == 0 else 16)
            run.font.color.rgb = FG if lvl == 0 else SUB


# 1 title
s = prs.slides.add_slide(BLANK); _bg(s)
tb = s.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.5), Inches(1.4))
p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = "VoxShield"
r.font.size = Pt(54); r.font.bold = True; r.font.color.rgb = FG
tb = s.shapes.add_textbox(Inches(0.9), Inches(3.4), Inches(11.5), Inches(1.0))
r = tb.text_frame.paragraphs[0].add_run()
r.text = "Voice Scam Shield for India's Investors — detect AI-cloned voices before money moves"
r.font.size = Pt(22); r.font.color.rgb = ACC
tb = s.shapes.add_textbox(Inches(0.9), Inches(4.6), Inches(11.5), Inches(1.6))
tf = tb.text_frame
for t in ["Track A · Digital Fraud & Scam Resilience", "23 Indic languages · narrowband telephony · on-prem · no advice, ever",
          "Team: Devansh Goenka et al.  ·  SANGYAN Investor Resilience Hackathon 2026"]:
    para = tf.add_paragraph(); run = para.add_run(); run.text = t
    run.font.size = Pt(15); run.font.color.rgb = SUB

add(prs.slides.add_slide(BLANK), "The problem: voice is the new attack surface",
    ["16+ crore Demat accounts; 70%+ of new retail accounts from Tier-2/3 cities",
     "Attackers use AI voice cloning / TTS to impersonate brokers, SEBI, family & 'finfluencers'",
     "Victims realise only after withdrawals are blocked — the money is already gone",
     "No accessible tool detects synthetic speech in Indic languages over phone-quality audio",
     (("deepfake \u201cfinfluencers\u201d are named explicitly in the challenge", 1))],
    kicker="Track A")

add(prs.slides.add_slide(BLANK), "Who we build for",
    ["A first-time investor in a Tier-2/3 city, more comfortable in a regional language",
     "Receives an urgent call/voice-note: \u201cSEBI se hoon, turant paise bhejo / share OTP\u201d",
     "Journey: hears clip \u2192 forwards to VoxShield \u2192 gets verdict + transcript + language \u2192 \u201cpause & verify\u201d via official channels",
     "Goal: recognise risk BEFORE a transfer \u2014 not after"])

add(prs.slides.add_slide(BLANK), "Solution: hear the human behind the call",
    ["Input: a suspicious call recording or voice note (any Indic language)",
     "Output: HUMAN / SYNTHETIC / ABSTAIN + transcript + detected language + reason codes",
     "Plain-language warning + \u201cPause & Verify\u201d action (SCORES / official broker / family)",
     "Works on narrowband phone audio (G.711 \u00b5/A-law, replay, packet-loss)",
     "Runs on-prem for privacy; never gives stock tips or recommendations"])

add(prs.slides.add_slide(BLANK), "Demo: one call, end to end",
    ["1  A genuine Hindi call \u2192 verdict HUMAN (language hi, transcript shown)",
     "2  A Bhashini-TTS scam call \u2192 verdict SYNTHETIC with reason codes",
     "3  Same clip through G.711 + replay \u2192 still flagged (behind the codec)",
     "4  Evidence disagrees \u2192 explicit ABSTAIN (honest uncertainty)",
     "5  Switch language (Tamil/Bengali) \u2192 works; on-prem, no cloud egress"])

add(prs.slides.add_slide(BLANK), "Architecture",
    ["Universal Input Profiler (codec/bandwidth/SNR) \u2192 adapter",
     "Fast L0 gate \u2192 SSL detector (XLS-R + RawBoost) + codec-aware branch",
     "Evidence brains: physics \u00b7 replay \u00b7 environment \u00b7 cross-codec \u00b7 speaker \u00b7 semantic\u2194prosody (ASR)",
     "Evidence arbitration \u2192 accept / ABSTAIN \u2192 VoxScore + novelty \u2192 verdict + reason codes",
     "Services: REST/WS API \u00b7 console \u00b7 IVR/SDK-ready"])

add(prs.slides.add_slide(BLANK), "Technology",
    ["Detector: fine-tuned XLS-R-300M + RawBoost, telephony augmentation",
     "Indic layer: Bhashini (ALD/ASR/TTS) + AIKosh corpora/models",
     "Speaker: ECAPA-TDNN verification + NeMo Sortformer diarization",
     "Robustness: PGD adversarial training + distilled edge student",
     "Benchmark: 12+ language Indic telephony; unseen-generator generalization"])

add(prs.slides.add_slide(BLANK), "Results (evaluated on Indic telephony)",
    ["Seen clean EER: 0.1\u20132%  (e.g., Odia clean 0.21%)",
     "G.711 \u00b5/A-law EER: 0.4\u20136%  (Odia G.711 0.40%)",
     "English-trained SOTA zero-shot on Indic: 33\u201393% EER \u2014 we cut it to single digits by fine-tuning on Indic",
     "Unseen-generator EER: 5\u201320% (external unseen generator) \u2014 the real generalization gap",
     "Abstention cuts genuine-caller false alarms 30\u201360% at 90% coverage; Cllr 0.3\u20131.0 bits"],
    kicker="Evidence")

add(prs.slides.add_slide(BLANK), "Bharat-first, privacy & guardrails",
    ["23 Indic languages; voice-first; low-bandwidth telephony; low-end edge model",
     "Privacy-by-design: on-prem / on-device inference; no SMS/OTP/PII harvesting",
     "No stock tips, no buy/sell/hold, no price prediction, no broker promotion",
     "Honest uncertainty: explicit ABSTAIN + confidence, never a false binary",
     "Public-good: investor-protection infrastructure, not a growth product"],
    kicker="Compliance")

add(prs.slides.add_slide(BLANK), "Impact, scalability & disclosure",
    ["Detects fraud before money moves; deployable by banks, brokers, SEBI/helplines, IVR, SDK",
     "Scales across languages, cheap phones and Tier-2/3 users",
     "Third-party components (disclosed): Bhashini APIs, AIKosh datasets/models, HuggingFace/NVIDIA models",
     "Licence note: IndicSynth is CC-BY-NC (research/benchmark use only)",
     "Roadmap: telephony gateway \u2192 voice-finfluencer content checks \u2192 district-language expansion"],
    kicker="Scale")

prs.save("SANGYAN_VoxShield.pptx")
print("wrote SANGYAN_VoxShield.pptx  (", len(prs.slides.__iter__.__self__._sldIdLst), "slides )")
