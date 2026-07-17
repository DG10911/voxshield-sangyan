#!/usr/bin/env python3
"""Build VoxShield investor/jury pitch deck (.pptx) from the presentation's facts.
All metrics are the measured/held-out numbers used in presentation.html."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- brand palette (matches presentation.html) -------------------------------
BG     = RGBColor(0x06, 0x0D, 0x18)   # deep navy
CARD   = RGBColor(0x0D, 0x1B, 0x30)
CARD2  = RGBColor(0x11, 0x21, 0x3A)
T1     = RGBColor(0xE8, 0xF0, 0xFF)   # primary text
T2     = RGBColor(0x93, 0xA9, 0xC8)   # secondary
T3     = RGBColor(0x64, 0x80, 0x9F)   # muted
ACC    = RGBColor(0x3B, 0x82, 0xF6)   # blue accent
ACC2   = RGBColor(0x8A, 0xB6, 0xFF)   # light blue
GRN    = RGBColor(0x22, 0xC5, 0x5E)
AMB    = RGBColor(0xF5, 0x9E, 0x0B)
RED    = RGBColor(0xEF, 0x44, 0x44)
LINE   = RGBColor(0x1E, 0x2C, 0x44)

FONT   = "Arial"
MONO   = "Consolas"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ---- helpers -----------------------------------------------------------------
def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element); s.shapes._spTree.insert(2, bg._element)
    return s

def _noline(sp): sp.line.fill.background(); sp.shadow.inherit = False

def rect(s, x, y, w, h, fill=CARD, line=None, radius=True):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line: shp.line.color.rgb = line; shp.line.width = Pt(1)
    else: shp.line.fill.background()
    shp.shadow.inherit = False
    if radius:
        try: shp.adjustments[0] = 0.06
        except Exception: pass
    return shp

def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp_after=4, line_sp=1.0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs[0], tuple): runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(sp_after); p.space_before = Pt(0); p.line_spacing = line_sp
        for (t, sz, col, bold, *rest) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col; r.font.bold = bold
            r.font.name = rest[0] if rest else FONT
    return tb

def kicker(s, text, x=0.7, y=0.55):
    rect(s, x, y, 0.14, 0.34, fill=ACC)
    txt(s, x+0.28, y-0.04, 8, 0.5, [[(text, 13, ACC2, True, MONO)]])

def title(s, lines, x=0.7, y=1.0, size=40, w=11.9):
    txt(s, x, y, w, 1.6, [[(seg[0], size, seg[1], True)] for seg in lines] if isinstance(lines[0], tuple) else
        [[(l, size, T1, True)] for l in lines])

def chip(s, x, y, w, text, col, tcol=None):
    c = rect(s, x, y, w, 0.42, fill=CARD, line=col)
    txt(s, x+0.12, y+0.03, w-0.2, 0.36, [[(text, 11.5, tcol or col, True, MONO)]], anchor=MSO_ANCHOR.MIDDLE)

def pagenum(s, n):
    txt(s, 12.3, 7.02, 1.0, 0.4, [[(f"{n:02d} / 14", 10, T3, False, MONO)]], align=PP_ALIGN.RIGHT)

def footnote(s, text, y=6.9):
    txt(s, 0.7, y, 11.9, 0.5, [[("▸  ", 10, ACC, False, MONO), (text, 10, T3, False, MONO)]])


# ============================================================ 1 · COVER
s = slide()
# faint accent band
rect(s, 0, 0, 13.333, 0.14, fill=ACC, radius=False)
txt(s, 0.7, 2.1, 12, 0.5, [[("UCO BANK · PUNJAB & SIND BANK · IIT KHARAGPUR  ·  PS2, AUDIO FORENSICS", 12, T2, True, MONO)]])
txt(s, 0.68, 2.55, 12, 2.0, [[("Vox", 88, T1, True), ("Shield", 88, ACC2, True)]])
txt(s, 0.7, 4.15, 11.5, 1.0,
    [[("Real-time AI voice-clone detection for India's voice-first banking.", 24, T2, False)],
     [("One layered, honest defence: ", 18, T2, False), ("Detection · Fairness · Liveness.", 18, ACC2, True)]],
    line_sp=1.15)
chip(s, 0.7, 5.7, 2.2, "● TRL-5 · WORKING SYSTEM", GRN)
chip(s, 3.05, 5.7, 3.0, "5-MODEL META-FUSION · 5.9% EER", ACC)
chip(s, 6.2, 5.7, 3.1, "FLAGS A FRAUD CALL IN 3.0 s", ACC2)
txt(s, 0.7, 6.55, 12, 0.5, [[("Team DigiSeva, SRM Institute of Science and Technology", 13, T3, True, MONO)]])
pagenum(s, 1)

# ============================================================ 2 · THE PROBLEM
s = slide(); kicker(s, "01 · THE THREAT")
title(s, [("Voice is becoming the easiest door into a bank account.", T1)], size=34, w=12)
cards = [
    ("◉  Cloning is cheap & instant", "Consumer TTS clones a speaker from seconds of public audio, a voicemail, a social clip. The fake sounds human.", AMB),
    ("☎  Telephony hides the evidence", "8 kHz G.711 lines smear the very artifacts most detectors rely on. Studio-trained models go blind on a real call.", AMB),
    ("⚖  English-centric models are unfair", "Detectors trained on English misread Indian languages as 'synthetic', genuine customers flagged as fraud.", RED),
    ("⧗  Damage happens in minutes", "Voice-authorised transfers and social-engineered agents act fast. Post-call forensics is not protection.", RED),
]
x0 = 0.7
for i, (h, b, col) in enumerate(cards):
    x = x0 + (i % 2) * 6.1; y = 2.15 + (i // 2) * 2.1
    rect(s, x, y, 5.8, 1.9, fill=CARD, line=LINE)
    txt(s, x+0.3, y+0.25, 5.2, 0.5, [[(h, 16, T1, True)]])
    txt(s, x+0.3, y+0.78, 5.2, 1.0, [[(b, 12.5, T2, False)]], line_sp=1.12)
footnote(s, "Industry data. Detection alone loses to a replay attack, a clone played through a speaker launders the digital artifacts. It needs layers.")
pagenum(s, 2)

# ============================================================ 3 · SOLUTION
s = slide(); kicker(s, "02 · OUR INNOVATION")
title(s, [("One layered, honest defence.", T1)], size=36)
txt(s, 0.7, 1.75, 12, 0.6, [[("VoxShield is not a model, it's defence-in-depth a bank can actually deploy and audit.", 16, T2, False)]])
sol = [
    ("🛡", "01 · DETECTION", "Five decorrelated detector families fused by a learned meta-stacker with calibrated, probability-true scores.", ACC),
    ("⚖", "02 · FAIRNESS", "Evaluated on genuine speech in 10 Indian languages, retrained Indic-aware, stress-tested under 8 kHz telephony.", GRN),
    ("◈", "03 · LIVENESS", "A random spoken-digit challenge a recording can't answer, content + liveness + timing verified together.", ACC2),
    ("✎", "04 · HONESTY", "Plain-language reason codes on every verdict · SHA-256 audit trail · never auto-blocks, routes to step-up.", AMB),
]
for i, (ic, h, b, col) in enumerate(sol):
    x = 0.7 + i * 3.05
    rect(s, x, 2.6, 2.85, 3.4, fill=CARD, line=LINE)
    ic_b = rect(s, x+0.3, 2.9, 0.75, 0.75, fill=CARD2, line=col)
    txt(s, x+0.3, 3.0, 0.75, 0.55, [[(ic, 22, col, False)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+0.3, 3.85, 2.3, 0.4, [[(h, 12.5, col, True, MONO)]])
    txt(s, x+0.3, 4.3, 2.35, 1.6, [[(b, 12, T2, False)]], line_sp=1.14)
footnote(s, "The same drop-in REST API guards every surface where a voice touches the bank, IVR, phone banking, agent-assist, KYC.")
pagenum(s, 3)

# ============================================================ 4 · FIVE MODELS (scoreboard)
s = slide(); kicker(s, "03 · FIVE AI MODELS, ONE BRAIN")
title(s, [("Weak alone. ", T1), ("Strong together.", ACC2)], size=36)
txt(s, 0.7, 1.75, 12, 0.5, [[("Each detector fails differently. A learned stacker trusts each exactly as much as it should, even down-weighting two of its own.", 14, T2, False)]])
# bars: name, eer%, width%, color
bars = [
    ("Wav2Vec2, alone", "62.4%", 0.96, RED),
    ("Acoustic DSP, alone", "46.9%", 0.72, RED),
    ("LFCC+CQCC head, alone", "26.8%", 0.41, T3),
    ("DistilHuBERT, alone", "23.0%", 0.35, T3),
    ("XLS-R 300M, best single", "16.0%", 0.25, T3),
    ("VoxShield meta-fusion", "5.9%", 0.091, GRN),
]
by = 2.5
for i, (name, val, w, col) in enumerate(bars):
    y = by + i * 0.66
    bold = col == GRN
    txt(s, 0.7, y, 3.4, 0.5, [[(name, 12.5, (GRN if bold else T2), bold)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 4.2, y+0.1, 6.6, 0.3, fill=CARD, line=None)
    rect(s, 4.2, y+0.1, max(0.15, 6.6*w), 0.3, fill=col, line=None)
    txt(s, 11.0, y, 1.5, 0.5, [[(val+" EER", 12.5, (GRN if bold else T1), bold, MONO)]], anchor=MSO_ANCHOR.MIDDLE)
footnote(s, "Equal-error-rate, lower is better · measured on held-out In-the-Wild data (n=320 disjoint test split). A 3× error reduction from fusion alone.")
pagenum(s, 4)

# ============================================================ 5 · MEASURED RESULTS
s = slide(); kicker(s, "04 · MEASURED RESULTS")
title(s, [("Numbers we can defend.", T1)], size=36)
txt(s, 0.7, 1.75, 12, 0.5, [[("Every figure is held-out, test clips the system never trained on. No cherry-picking, no projections dressed as measurements.", 14, T2, False)]])
stats = [
    ("5.9%", "Equal-error-rate", "down from 20.3% naive blend · held-out In-the-Wild", GRN),
    ("0.983", "ROC AUC", "near-perfect fake-vs-genuine ranking on unseen data", ACC2),
    ("93.4%", "Accuracy @0.70", "precision 93.5% · recall 92.9%", ACC2),
    ("3.0 s", "Time-to-flag", "synthetic call flagged in streaming, well inside 10 s", GRN),
    ("36.3→6.3%", "Fairness FP rate", "genuine Indic speech, Indic-aware + channel-aware", GRN),
    ("5+1", "Fused detectors", "one learned meta-stacker · calibrated output", ACC2),
]
for i, (v, k, b, col) in enumerate(stats):
    x = 0.7 + (i % 3) * 4.03; y = 2.55 + (i // 3) * 1.95
    rect(s, x, y, 3.8, 1.75, fill=CARD, line=LINE)
    rect(s, x, y, 1.5, 0.05, fill=ACC, radius=False)
    txt(s, x+0.3, y+0.2, 3.3, 0.7, [[(v, 30, col, True, MONO)]])
    txt(s, x+0.3, y+0.85, 3.3, 0.4, [[(k, 13, T1, True)]])
    txt(s, x+0.3, y+1.2, 3.3, 0.5, [[(b, 10.5, T2, False)]], line_sp=1.05)
footnote(s, "Protocol: neural models frozen; meta-stacker trained on 480 clips, evaluated on a disjoint 320-clip split, honest cross-dataset numbers.")
pagenum(s, 5)

# ============================================================ 6 · FAIRNESS
s = slide(); kicker(s, "05 · THE FAIRNESS BREAKTHROUGH")
title(s, [("A detector that flags honest customers is a failed detector.", T1)], size=30, w=12)
# morph number
rect(s, 0.7, 2.4, 12, 2.0, fill=CARD, line=LINE)
txt(s, 1.2, 2.9, 5, 1.2, [[("36.3%", 44, T3, True, MONO)]], anchor=MSO_ANCHOR.MIDDLE)
txt(s, 4.4, 2.9, 1.4, 1.2, [[("→", 40, ACC, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, 5.6, 2.75, 6, 1.4, [[("6.3%", 60, GRN, True, MONO)]], anchor=MSO_ANCHOR.MIDDLE)
txt(s, 1.2, 3.95, 11, 0.5, [[("False-positive rate on genuine Indian-language speech, before → after Indic-aware retraining.", 13, T2, False)]])
txt(s, 0.7, 4.7, 12, 0.9,
    [[("Evaluated across 10 Indian languages (Hindi, Tamil, Bengali, Telugu, Marathi, Kannada, Malayalam, Gujarati, Punjabi, Assamese), ", 13.5, T2, False)],
     [("and we publish the uncomfortable numbers too. This measured fairness is the moat a public-sector bank can actually procure.", 13.5, T2, False)]], line_sp=1.15)
footnote(s, "Internal benchmarking · IndicVoices genuine speech · 30 clips per language · threshold 0.70.")
pagenum(s, 6)

# ============================================================ 7 · LIVENESS + ARREST SHIELD
s = slide(); kicker(s, "06 · BEYOND DETECTION")
title(s, [("Liveness & the Digital Arrest Shield.", T1)], size=32)
# left: liveness
rect(s, 0.7, 2.15, 5.8, 4.0, fill=CARD, line=LINE)
txt(s, 1.0, 2.4, 5.2, 0.5, [[("◈  LIVENESS, defeats replay", 15, ACC2, True)]])
for i, (h, b) in enumerate([
    ("Content", "Whisper ASR, do the spoken digits match the random challenge?"),
    ("Liveness", "The full ensemble scores the reply, a live human, not a synthesis engine?"),
    ("Timing", "Did the answer arrive in the challenge window? A pass needs all three."),
]):
    y = 3.0 + i*0.95
    rect(s, 1.0, y, 0.42, 0.42, fill=CARD2, line=GRN)
    txt(s, 1.0, y, 0.42, 0.42, [[("✓", 14, GRN, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 1.6, y-0.05, 4.7, 0.4, [[(h, 13.5, T1, True)]])
    txt(s, 1.6, y+0.32, 4.7, 0.6, [[(b, 11, T2, False)]], line_sp=1.05)
# right: arrest shield
rect(s, 6.75, 2.15, 5.85, 4.0, fill=CARD, line=LINE)
txt(s, 7.05, 2.4, 5.3, 0.5, [[("🛡  DIGITAL ARREST SHIELD", 15, AMB, True)]])
txt(s, 7.05, 2.95, 5.3, 1.5,
    [[("Fuses the 5-model clone score with scam-intent read from the transcript, ", 12.5, T2, False)],
     [("catches a cloned/social-engineered voice running the ", 12.5, T2, False), ("digital-arrest playbook", 12.5, AMB, True), (".", 12.5, T2, False)]], line_sp=1.15)
rect(s, 7.05, 4.35, 5.3, 1.55, fill=CARD2, line=None)
txt(s, 7.3, 4.5, 5.0, 1.4,
    [[("“Officer Sharma, CBI Cybercrime…”", 12, T2, False, MONO)],
     [("→ verdict CRITICAL", 13, RED, True, MONO)],
     [("→ clone 0.95 · scam-intent 1.0 · 14 flags", 11.5, T3, False, MONO)],
     [("→ routed to step-up · SHA-256 audit", 11.5, GRN, False, MONO)]], line_sp=1.2)
footnote(s, "Multilingual (Hindi / Hinglish / English) transcription. No competitor combines clone-detection with scam-intent on the live call.")
pagenum(s, 7)

# ============================================================ 8 · ARCHITECTURE
s = slide(); kicker(s, "07 · REAL PRODUCT, SYSTEM ARCHITECTURE")
title(s, [("Not a deck. A working system.", T1)], size=32)
layers = [
    ("CHANNELS", "Operator console · Live 2-party call · Liveness · Bank demo · IVR / 8 kHz G.711", ACC),
    ("API, FastAPI", "/analyze · /stream-analyze · /analyze-call · /liveness · /tts · WS /ws/call · /docs", ACC2),
    ("CORE ENGINE", "Sliding windows → 89-dim feature bank → 5 detectors → meta-fusion → reason codes", GRN),
    ("TRUST & DATA", "SHA-256 audit ledger · action log · step-up routing · auditable JSON artifacts", AMB),
    ("OPS & SCALE", "Docker (models baked in) · stateless · GPU sub-sec / CPU 3-8 s · on-prem / air-gapped", T2),
]
for i, (h, b, col) in enumerate(layers):
    y = 2.15 + i*0.92
    rect(s, 0.7, y, 2.5, 0.78, fill=CARD2, line=col)
    txt(s, 0.85, y, 2.3, 0.78, [[(h, 12.5, col, True, MONO)]], anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 3.4, y, 9.2, 0.78, fill=CARD, line=LINE)
    txt(s, 3.65, y, 8.8, 0.78, [[(b, 12.5, T2, False)]], anchor=MSO_ANCHOR.MIDDLE)
    if i < 4:
        txt(s, 1.8, y+0.72, 0.4, 0.25, [[("▼", 9, ACC, False)]], align=PP_ALIGN.CENTER)
pagenum(s, 8)

# ============================================================ 9 · WHY US (moat)
s = slide(); kicker(s, "08 · THE MOAT")
title(s, [("What a single-model API can't do.", T1)], size=32)
rows = [
    ("Decorrelated multi-model detection", "✕ one model, one blind spot", "✓ five families + learned stacker"),
    ("Calibrated probabilities", "✕ raw scores", "✓ Platt-calibrated"),
    ("Explainable verdicts", "✕ black box", "✓ SSL · PH · HF · PR · BR reason codes"),
    ("Indian-language fairness, measured", "✕ English-centric", "✓ 10 languages, FP rates published"),
    ("Replay-attack defence", "✕ detection only", "✓ challenge-response liveness"),
    ("Verdict inside the live call", ", upload after the fact", "✓ streaming, flagged at 3.0 s"),
    ("On-prem / air-gapped", "✕ cloud only", "✓ models baked into Docker"),
    ("Customer-safety guarantee", ", score, then your problem", "✓ never auto-block, step-up"),
]
# header
hy = 2.1
txt(s, 0.7, hy, 5.0, 0.4, [[("CAPABILITY", 11, T3, True, MONO)]])
txt(s, 5.9, hy, 3.2, 0.4, [[("TYPICAL API", 11, T3, True, MONO)]])
txt(s, 9.3, hy, 3.4, 0.4, [[("VOXSHIELD", 11, ACC2, True, MONO)]])
for i, (c, a, b) in enumerate(rows):
    y = 2.55 + i*0.53
    if i % 2 == 0: rect(s, 0.7, y-0.03, 11.95, 0.5, fill=CARD, line=None)
    txt(s, 0.85, y, 5.0, 0.45, [[(c, 11.5, T1, False)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 5.9, y, 3.3, 0.45, [[(a, 11, T3, False)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 9.3, y, 3.4, 0.45, [[(b, 11, GRN, False)]], anchor=MSO_ANCHOR.MIDDLE)
footnote(s, "Architecture comparison against the common commercial pattern, a single neural detector behind a cloud endpoint. Not a benchmark of any vendor.")
pagenum(s, 9)

# ============================================================ 10 · MARKET
s = slide(); kicker(s, "09 · THE MARKET")
title(s, [("India's banking future is ", T1), ("voice-first.", ACC2)], size=36)
txt(s, 0.7, 1.9, 12, 1.0,
    [[("Phone banking, IVR and voice-led payments reach millions that apps never will, across languages, literacy", 15, T2, False)],
     [("levels and feature phones. Every one of those calls is a voice-cloning attack surface that must be trusted.", 15, T2, False)]], line_sp=1.15)
mk = [
    ("◎  The channel banks can't abandon", "Voice is the default banking interface for millions, and the first channel cloning attacks target. Table stakes for every call centre."),
    ("⚖  Fairness is the moat", "Anyone can wrap an English deepfake model in an API. Measured fairness across 10 Indian languages is what a public-sector bank can procure."),
    ("⬡  Beyond banking", "Same engine: government helplines, telecom anti-vishing, KYC re-verification, courtroom audio forensics, media verification."),
]
for i, (h, b) in enumerate(mk):
    x = 0.7 + i*4.03
    rect(s, x, 3.2, 3.8, 2.7, fill=CARD, line=LINE)
    txt(s, x+0.3, 3.45, 3.3, 0.7, [[(h, 14, T1, True)]], line_sp=1.05)
    txt(s, x+0.3, 4.35, 3.3, 1.4, [[(b, 12, T2, False)]], line_sp=1.15)
pagenum(s, 10)

# ============================================================ 11 · TRACTION / DEPLOY
s = slide(); kicker(s, "10 · DEPLOYMENT & TRUST")
title(s, [("Built for a bank's basement, not our cloud.", T1)], size=30)
tl = [
    ("#", "SHA-256 audit trail", "Every clip hashed for the ledger, raw audio never stored in clear.", ACC2),
    ("⊘", "Never auto-block", "A layered fraud signal, not a judge. No customer locked out by a model score.", GRN),
    ("✎", "Explainable by default", "Reason codes on every verdict + /why-fusion plain-language account.", ACC2),
    ("◫", "Accessible & compliant", "GOV.UK / USWDS patterns, WCAG 2.1 AA, risk by icon + text, not colour.", AMB),
]
for i, (ic, h, b, col) in enumerate(tl):
    x = 0.7 + (i % 2)*6.1; y = 2.3 + (i // 2)*1.7
    rect(s, x, y, 5.8, 1.5, fill=CARD, line=LINE)
    rect(s, x+0.28, y+0.28, 0.7, 0.7, fill=CARD2, line=col)
    txt(s, x+0.28, y+0.28, 0.7, 0.7, [[(ic, 20, col, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+1.2, y+0.3, 4.4, 0.4, [[(h, 14, T1, True)]])
    txt(s, x+1.2, y+0.72, 4.4, 0.7, [[(b, 11.5, T2, False)]], line_sp=1.1)
# deploy options strip
rect(s, 0.7, 5.85, 11.95, 0.95, fill=CARD2, line=None)
txt(s, 0.95, 5.98, 11.5, 0.7,
    [[("Deploy: ", 12, T3, True, MONO), ("On-prem GPU ", 12, GRN, True, MONO), ("sub-second · ", 12, T2, False, MONO),
      ("Cloud CPU ", 12, ACC2, True, MONO), ("3-8 s · ", 12, T2, False, MONO),
      ("Free 16 GB container ", 12, ACC2, True, MONO), (", evaluation. Docker bakes all models in; no runtime downloads.", 12, T2, False, MONO)]],
    anchor=MSO_ANCHOR.MIDDLE)
pagenum(s, 11)

# ============================================================ 12 · ROADMAP
s = slide(); kicker(s, "11 · ROADMAP, TRL 5 → 9")
title(s, [("The generators improve. So does the shield.", T1)], size=30)
# TRL meter
mx = 0.7
for i in range(9):
    filled = i < 5
    rect(s, mx + i*1.32, 2.1, 1.15, 0.22, fill=(ACC if filled else CARD), line=None)
    txt(s, mx + i*1.32, 2.34, 1.15, 0.3, [[(str(i+1), 9, (ACC2 if filled else T3), False, MONO)]], align=PP_ALIGN.CENTER)
txt(s, 0.7, 2.7, 12, 0.35, [[("TRL-5 today (validated working system) → target TRL-9 (national production).", 12, T2, False, MONO)]])
ph = [
    ("SHIPPED · TRL-5", "5-model fusion 5.9% EER · fairness 10 langs · liveness · live call · Docker + audit", GRN),
    ("0-3 mo → TRL-6", "ONNX-quantised AASIST-L for CPU streaming · bank-sandbox pilot on recorded IVR", ACC2),
    ("3-6 mo → TRL-7", "Indic fine-tuning (Telugu first) · continual learning · live-traffic shadow deploy", ACC2),
    ("6-12 mo → TRL-8", "Source attribution · adversarial-robust training · regulator-ready audit pack", ACC2),
    ("12 mo+ → TRL-9", "Multi-bank rollout · toward all 22 languages · shared fraud-signal exchange", ACC),
]
for i, (h, b, col) in enumerate(ph):
    y = 3.35 + i*0.68
    rect(s, 0.7, y, 2.7, 0.56, fill=CARD2, line=col)
    txt(s, 0.85, y, 2.5, 0.56, [[(h, 11, col, True, MONO)]], anchor=MSO_ANCHOR.MIDDLE)
    txt(s, 3.6, y, 9.0, 0.56, [[(b, 12, T2, False)]], anchor=MSO_ANCHOR.MIDDLE)
footnote(s, "TRL-5 is measured. Phases 1-4 are engineering plans with target windows, nothing beyond TRL-5 is claimed as done.")
pagenum(s, 12)

# ============================================================ 13 · THE ASK
s = slide(); kicker(s, "12 · THE ASK")
title(s, [("Fund the pilot. Harden for India.", T1)], size=34)
ask = [
    ("Deploy", "Bank-sandbox pilot on recorded IVR traffic, prove the false-positive budget on real calls."),
    ("Harden", "Indic fine-tuning + continual learning to push fairness further and track new generators."),
    ("Certify", "Regulator-ready audit pack: model cards, published fairness, decision logs → procurement."),
]
for i, (h, b) in enumerate(ask):
    x = 0.7 + i*4.03
    rect(s, x, 2.3, 3.8, 2.5, fill=CARD, line=LINE)
    rect(s, x, 2.3, 1.4, 0.05, fill=ACC, radius=False)
    txt(s, x+0.3, 2.55, 3.3, 0.5, [[(f"0{i+1}", 26, ACC2, True, MONO)]])
    txt(s, x+0.3, 3.15, 3.3, 0.4, [[(h, 16, T1, True)]])
    txt(s, x+0.3, 3.6, 3.3, 1.1, [[(b, 12, T2, False)]], line_sp=1.15)
rect(s, 0.7, 5.2, 11.95, 1.2, fill=CARD2, line=ACC)
txt(s, 1.0, 5.4, 11.4, 0.9,
    [[("Why now: ", 15, ACC2, True), ("a working TRL-5 system, an honest fairness moat, and a voice-first banking market", 15, T1, False)],
     [("that is under active attack today. Detection, fairness and liveness, one honest defence, ready to pilot.", 15, T1, False)]], line_sp=1.2)
pagenum(s, 13)

# ============================================================ 14 · CLOSE
s = slide()
rect(s, 0, 0, 13.333, 0.14, fill=ACC, radius=False)
txt(s, 0.7, 2.2, 12, 0.5, [[("THANK YOU", 13, ACC, True, MONO)]])
txt(s, 0.68, 2.8, 12.2, 1.6,
    [[("Detection. ", 46, ACC2, True), ("Fairness. ", 46, GRN, True), ("Liveness.", 46, AMB, True)]])
txt(s, 0.7, 3.9, 12, 0.8, [[("One layered, honest defence for India's voice-first banking future.", 22, T2, False)]])
txt(s, 0.7, 5.4, 12, 1.0,
    [[("VoxShield, Team DigiSeva · SRM Institute of Science and Technology", 15, T1, True)],
     [("UCO Bank · Punjab & Sind Bank · IIT Kharagpur, PS2 · Audio Forensics for Voice Security · TRL-5", 12, T3, False, MONO)]], line_sp=1.4)
pagenum(s, 14)

out = "/Users/devanshgoenka/conductor/workspaces/voxshield/san-antonio/VoxShield_Pitch.pptx"
prs.save(out)
print("saved", out, "·", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
