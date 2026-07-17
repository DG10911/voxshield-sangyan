#!/usr/bin/env python3
"""Build VoxShield_DGX_A100_Proposal.docx, a formal GPU-access proposal for the HOD / AI-DL Lab."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0B, 0x1F, 0x33)
BLUE = RGBColor(0x1E, 0x5F, 0xB0)
STEEL= RGBColor(0x33, 0x4E, 0x68)
GREY = RGBColor(0x5A, 0x6B, 0x7B)
DARK = RGBColor(0x1E, 0x1E, 0x1E)
GREEN= RGBColor(0x1B, 0x7A, 0x43)

doc = Document()
base = doc.styles["Normal"]; base.font.name = "Calibri"; base.font.size = Pt(10.5); base.font.color.rgb = DARK
base.paragraph_format.space_after = Pt(5); base.paragraph_format.line_spacing = 1.12
sec = doc.sections[0]
sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.8); sec.left_margin = Inches(0.9); sec.right_margin = Inches(0.9)

def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexc); tcPr.append(sh)

def para(text="", size=10.5, color=DARK, bold=False, italic=False, before=0, after=5, align=None, font="Calibri"):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    if align is not None: p.alignment = align
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = color; r.bold = bold; r.italic = italic; r.font.name = font
    return p

def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = DARK
    r = p.add_run(text); r.font.size = Pt(10.5); r.font.color.rgb = DARK
    return p

def h(num, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{num}.  "); r.font.size = Pt(13); r.bold = True; r.font.color.rgb = BLUE
    r = p.add_run(text); r.font.size = Pt(13); r.bold = True; r.font.color.rgb = NAVY
    # bottom rule
    pPr = p._p.get_or_add_pPr(); pb = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "3"); b.set(qn("w:color"), "C7D6E6")
    pb.append(b); pPr.append(pb)
    return p

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    hd = t.rows[0].cells
    for i, ht in enumerate(headers):
        shade(hd[i], "0B1F33"); pr = hd[i].paragraphs[0]; pr.paragraph_format.space_after = Pt(1)
        run = pr.add_run(ht); run.bold = True; run.font.size = Pt(9.5); run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            if ri % 2 == 1: shade(cells[i], "EEF3F9")
            pr = cells[i].paragraphs[0]; pr.paragraph_format.space_after = Pt(1)
            run = pr.add_run(str(val)); run.font.size = Pt(9)
            if i == 0: run.bold = True
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

# ─────────────────────────── TITLE BLOCK ───────────────────────────
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
r = p.add_run("REQUEST FOR NVIDIA DGX A100 ACCESS"); r.bold = True; r.font.size = Pt(19); r.font.color.rgb = NAVY
para("VoxShield: Real-time AI Voice-Clone Detection for Indian Banking", 13.5, BLUE, bold=True, after=6)
# sub-line box
tb = doc.add_table(rows=1, cols=1); tb.style = "Table Grid"; shade(tb.rows[0].cells[0], "F2F6FB")
c = tb.rows[0].cells[0]; cp = c.paragraphs[0]; cp.paragraph_format.space_after = Pt(2)
for txt, bold in [("Team DigiSeva  ·  SRM Institute of Science & Technology\n", True),
                  ("PSB Hackathon (UCO Bank · Punjab & Sind Bank · IIT Kharagpur), Problem Statement PS2\n", False),
                  ("Also positioned for Smart India Hackathon (SIH) and productisation.", False)]:
    run = cp.add_run(txt); run.font.size = Pt(10); run.bold = bold; run.font.color.rgb = STEEL
doc.add_paragraph().paragraph_format.space_after = Pt(2)

para("To:  The Head of Department [Dept] · AI/DL Lab In-charge · Faculty Mentor [Name]      |      Date: [DD Month 2026]",
     9.5, GREY, italic=True, after=6)

# ─────────────────────────── 1. EXECUTIVE SUMMARY ───────────────────────────
h("1", "Executive Summary")
para("We respectfully request access to the institute's NVIDIA DGX A100 to train and fine-tune the models behind VoxShield, a real-time AI voice-clone detection system for Indian banking. VoxShield is already a working system (TRL-5); the requested compute will let us scale it from ten to all twenty-two scheduled Indian languages, harden it against adversarial attacks, and prepare a certified, reproducible build for the PSB Hackathon, Smart India Hackathon, and subsequent productisation. We seek a modest, scheduled allocation on the shared DGX partition over three months, under faculty mentorship, using only public datasets.")

# ─────────────────────────── 2. THE PROJECT ───────────────────────────
h("2", "The Project, VoxShield (a working system, TRL-5)")
para("VoxShield is a defence-in-depth architecture, not a single model: five decorrelated detectors fused by a learned meta-stacker, a fairness layer measured on Indian languages, and a challenge-response liveness check. All figures below are measured on held-out data.", after=4)
bullet("5.9% equal-error-rate on held-out In-the-Wild data, 0.983 ROC-AUC, ~94% accuracy, approximately 3x better than the best single model (XLS-R 300M at 16.0% EER).", "Detection: ")
bullet("false-positives on genuine Indian-language speech reduced from 36.3% (English-only) to 6.3% via Indic-aware retraining and channel-aware thresholds, measured across ten Indian languages (IndicVoices).", "Fairness: ")
bullet("Indian-language voice-clone recall lifted from 42% to 82% with a language-routed Indic-aware detector.", "Indic recall: ")
bullet("a random spoken-digit challenge (content + liveness + timing) that defeats replay attacks.", "Liveness: ")
bullet("fuses the clone score with scam-intent read from a Hindi / Hinglish / English transcript.", "Digital Arrest Shield: ")
bullet("flags a fraud call in 3.0 seconds, well inside a 10-second target.", "Streaming: ")
bullet("on-premise / air-gapped (Docker, models baked in), SHA-256 audit trail, and it never auto-blocks, high risk routes to step-up verification.", "Deployment: ")
para("Stack: Python, FastAPI, PyTorch, HuggingFace Transformers (wav2vec2 / XLS-R 300M / DistilHuBERT), Whisper and faster-whisper ASR, ECAPA speaker embeddings, and Silero VAD. Trained on ASVspoof 2019 LA, WaveFake, CodecFake, In-the-Wild, IndicVoices, and synthesised MMS-TTS / VITS Indic deepfakes.",
     9.5, STEEL, italic=True, before=2)

# ─────────────────────────── 3. PRODUCT & FUTURE POTENTIAL ───────────────────────────
h("3", "Product & Future Potential")
para("VoxShield is positioned as a deployable product, not a class project.", after=4)
bullet("banking fraud desks, phone banking, IVR, UPI / voice-authorised payments, KYC. The same drop-in REST API extends to government helplines, telecom anti-vishing, insurance, healthcare, legal forensics and contact centres, twelve sectors in all.", "Beachhead & reach: ")
bullet("enterprise SaaS, usage-based API, on-premise licensing, government and banking contracts, and OEM / white-label.", "Business model: ")
bullet("India's voice-first banking reaches hundreds of millions across languages and feature phones, and voice-cloning fraud such as digital-arrest scams is a growing threat. This aligns directly with Smart India Hackathon themes in cybersecurity, fintech and Digital India.", "National relevance: ")
bullet("fine-tuning across all twenty-two scheduled Indian languages, continual learning for new voice generators, source attribution, and adversarial-robust training (Malafide / Malacopula-class attacks), each requiring the DGX.", "Roadmap that needs the DGX: ")

# ─────────────────────────── 4. WHY THE DGX A100 ───────────────────────────
h("4", "Why the NVIDIA DGX A100, Planned Workloads")
para("The DGX A100 (8x A100 Tensor-Core GPUs, 320 GB total VRAM, NVIDIA Ampere) is the right tool for four GPU-bound workloads:", after=4)
bullet("fine-tuning XLS-R 300M and AASIST-L deepfake detectors across ten to twenty-two Indian languages, using large-batch, mixed-precision, multi-GPU data-parallel training.", "1)  ")
bullet("large-scale synthetic Indic deepfake generation (VITS / MMS-TTS), thousands of clips in parallel to build a balanced training corpus.", "2)  ")
bullet("adversarial-robust training and continual-learning experiments that are GPU-heavy and impractical on rented single GPUs.", "3)  ")
bullet("retraining the full five-model ensemble and meta-stacker with honest held-out cross-dataset evaluation, in hours instead of days.", "4)  ")
para("Current constraint: the team presently rents single cloud GPUs and runs CPU inference on an 8 GB machine, which bottlenecks iteration ahead of the hackathon deadline.",
     9.5, STEEL, italic=True, before=2)

# ─────────────────────────── 5. RESOURCE REQUEST ───────────────────────────
h("5", "Resource Request")
para("We make a modest, specific request in keeping with shared-lab etiquette:", after=4)
table(["Item", "Requested"],
      [["Compute", "2-4x NVIDIA A100 GPUs (40 GB or 80 GB), DGX A100 shared partition"],
       ["Duration", "3 months (phased plan in Section 6)"],
       ["Storage", "~500 GB scratch space for datasets and model checkpoints"],
       ["Access method", "Key-based SSH to the DGX node, or Slurm batch partition / JupyterHub, whichever the lab prefers"],
       ["Software", "PyTorch, CUDA, HuggingFace; containerised via the lab's NGC / Docker images"]],
      widths=[1.5, 5.2])
para("Server / hostname, SSH endpoint and account provisioning as assigned by the lab in-charge, we will supply our public SSH key and abide by all lab usage and scheduling policies.",
     9.5, DARK, bold=True)

# ─────────────────────────── 6. TIMELINE ───────────────────────────
h("6", "Three-Month Plan")
table(["Phase", "Activities", "Deliverable"],
      [["Month 1", "Environment setup, dataset staging, baseline reproduction, Indic fine-tune (10 languages)",
        "Reproduced 5.9% EER and improved Indic recall"],
       ["Month 2", "Scale to all 22 languages, synthetic Indic deepfake corpus, channel-aware fairness re-runs",
        "Published per-language fairness report"],
       ["Month 3", "Adversarial-robust training, source attribution, final held-out evaluation, hackathon / SIH build + model card",
        "Certified, reproducible results"]],
      widths=[0.9, 3.9, 1.9])

# ─────────────────────────── 7. DATA, SECURITY & GOVERNANCE ───────────────────────────
h("7", "Data, Security & Governance")
para("VoxShield uses only public datasets. Raw audio is SHA-256 hashed and no personal data is stored. All work will comply with lab usage policy and scheduling. The institution and the AI/DL Lab will be acknowledged in our hackathon submission and in any resulting publication.")

# ─────────────────────────── 8. EXPECTED OUTCOMES & REQUEST ───────────────────────────
h("8", "Expected Outcomes & Request")
bullet("a hackathon- and SIH-ready, reproducible, deployable system.", "")
bullet("a possible research publication on Indic anti-spoofing and fairness.", "")
bullet("a reusable Indian-language anti-spoofing asset that remains with the lab.", "")
para("We would be grateful if the Head of Department could guide the process and connect our team with the AI/DL Lab in-charge, under faculty mentorship. We are happy to present the working system on request and to adhere to any conditions the lab specifies.",
     before=3)

# ─────────────────────────── SIGNATURE BLOCK ───────────────────────────
doc.add_paragraph().paragraph_format.space_after = Pt(6)
para("Respectfully submitted,", 10.5, DARK, before=4, after=8)
table(["Team DigiSeva, Member", "Reg. No.", "Department / Year", "Signature"],
      [["[Name 1]", "[Reg. No.]", "[Dept / Year]", ""],
       ["[Name 2]", "[Reg. No.]", "[Dept / Year]", ""],
       ["[Name 3]", "[Reg. No.]", "[Dept / Year]", ""],
       ["[Name 4]", "[Reg. No.]", "[Dept / Year]", ""]],
      widths=[2.2, 1.4, 1.9, 1.2])
para("Faculty Guide: [Name, Designation] ____________________________        Date: [DD Month 2026]",
     10, DARK, before=4)

out = "/Users/devanshgoenka/conductor/workspaces/voxshield/san-antonio/VoxShield_DGX_A100_Proposal.docx"
doc.save(out); print("saved", out)
